import os
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel

load_dotenv()
my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("there is no API key")

client = Groq(api_key=my_api_key)

model= "openai/gpt-oss-120b"
job_description = """
Description: 
Description
Do you want to solve real customer problems through innovative technology? Do you enjoy working on scalable services in a collaborative team environment? Do you want to see your code directly impact millions of customers worldwide?

At Amazon, we hire the best minds in technology to innovate and build on behalf of our customers. Customer obsession is part of our company DNA, which has made us one of the world's most beloved brands.

Our Software Development Engineers (SDEs) use modern technology to solve complex problems while seeing their work's impact first-hand. The challenges SDEs solve at Amazon are meaningful and influence millions of customers, sellers, and products globally. We seek individuals passionate about creating new products, features, and services while managing ambiguity in an environment where development cycles are measured in weeks, not years.

At Amazon, we believe in ownership at every level. As an SDE-I, you'll own the entire lifecycle of your code - from design through deployment and ongoing operations. This ownership mindset, combined with our commitment to operational excellence, ensures we deliver the highest quality solutions for our customers.

We're looking for curious minds who think big and want to define tomorrow's technology. At Amazon, you'll grow into the high-impact engineer you know you can be, supported by a culture of learning and mentorship. Every day brings exciting new challenges and opportunities for personal growth.

We will take your location preferences into consideration. Preferences are based on business availability and are not guaranteed. Applicants will be considered at all locations in the United States including but not limited to:


During your application, you will have the opportunity to highlight your expertise in one or more of these specialized areas:

- Machine Learning • Distributed Systems and Data Management • Database Systems • Quantum Computing • Network Development • Query Processing and Optimization • Automated Reasoning • Embedded Systems • Data Engineering • Mobile Development • Game Development

While team placement is based on business needs, sharing your specific skills and interests helps us better align your experience with potential roles.

Applications are reviewed on a rolling basis. For an update on your status or to confirm your application was submitted successfully, please login to your candidate portal at Amazon.jobs. Amazon works with a high volume of applicants so we appreciate your patience as we review applications.


Key job responsibilities
- Collaborate and communicate effectively with experienced cross-disciplinary Amazonians to design, build, and operate innovative products and services that delight our customers, while participating in technical discussions to drive solutions forward.
- Design and develop scalable solutions using cloud-native architectures and microservices in a large distributed computing environment.
- Participate in code reviews and contribute to technical documentation.
- Build and maintain resilient distributed systems that are scalable, fault-tolerant, and cost-effective.
- Leverage and contribute to the development of GenAI and AI-powered tools to enhance development productivity while staying current with emerging technologies.
- Write clean, maintainable code following best practices and design patterns.
- Work in an agile environment practicing CI/CD principles while participating in operational responsibilities including on-call duties.
- Demonstrate operational excellence through monitoring, troubleshooting, and resolving production issues.

Basic Qualifications
- Experience with at least one general-purpose programming language such as Java, Python, C++, C#, Go, Rust, or TypeScript
- Experience with data structure implementation, basic algorithm development, and/or object-oriented design principles
- Currently has, or is in the process of obtaining a bachelor’s degree in Computer Science, Computer Engineering, Data Science, Information Systems, or related STEM fields
- Must be 18 years of age of older

Preferred Qualifications
- Experience from previous technical internship(s) or demonstrated project experience
- Experience with one or more of the following: AI tools for development productivity, Cloud platforms (preferably AWS), Database systems (SQL and NoSQL), Contributing to open-source projects, Version control systems, Debugging and troubleshooting complex systems
- Demonstrated ability to learn and adapt to new technologies quickly
- Basic understanding of software development lifecycle (SDLC)
- Strong problem-solving and analytical skills
- Excellent written and verbal communication skills

Amazon is an equal opportunity employer and does not discriminate on the basis of protected veteran status, disability, or other legally protected status.

Our inclusive culture empowers Amazonians to deliver the best results for our customers. If you have a disability and need a workplace accommodation or adjustment during the application and hiring process, including support for the interview or onboarding process, please visit https://amazon.jobs/content/en/how-we-hire/accommodations for more information. If the country/region you’re applying in isn’t listed, please contact your Recruiting Partner.


The base salary range for this position is listed below. Your Amazon package will include sign-on payments and restricted stock units (RSUs). Final compensation will be determined based on factors including experience, qualifications, and location. Amazon also offers comprehensive benefits including health insurance (medical, dental, vision, prescription, Basic Life & AD&D insurance and option for Supplemental life plans, EAP, Mental Health Support, Medical Advice Line, Flexible Spending Accounts, Adoption and Surrogacy Reimbursement coverage), 401(k) matching, paid time off, and parental leave. Learn more about our benefits at https://amazon.jobs/en/benefits.

"""
class JobDescription(BaseModel):
    role:str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: list[str] | None
    educational_requirement: list[str]
    resposibilities: list[str] 

jobd_schema = JobDescription.model_json_schema()

system_prompt = f"""
you are na expert hr assistant

your job is to analyze job description and extract
structure information from them.

return only json matching this schema:
{jobd_schema}
Important
do not return the schema itslf
do not return field like "propties", "titlee", or "type"
fill th schema with actual information extracted from the job description

If minimu experienc e is not mentioned, return null.
if information is not mentioned, return empty list.
do not invent any infromation, return only what is present in the job description.
"""
user_prompt = f"""
Analyze th following job description and extract structured information from it.
{job_description}"""

messages_system = {
    "role" : "system",
    "content": system_prompt

}
message_user  = {
    "role" : "user",
    "content" : user_prompt
}
response_format = {
    "type": "json_object",   # FIX: "json_schema" needs a schema attached, json_object works
}

messages = [messages_system, message_user]

response = client.chat.completions.create(
    model = model
    , messages = messages,
    response_format = response_format

)
answer = response.choices[0].message.content

raw_json = answer

import json 
job_data = json.loads(raw_json)
job = JobDescription(**job_data)

print(job.minimum_experience)
print(job.educational_requirement)
print(job.resposibilities)   # FIX: matches the field name in your class

class MatchResult(BaseModel):
    match_score : float
    details: dict

class Expereince(BaseModel):
    company: str | None
    role: str | None
    duration: str | None
    description: str | None
    skills_used : list[str] = []

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    experience:  float | None = None  

    skills: list[str] = []
    experiences: list[Expereince] = []
    eduation: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []



resume_schema = Resume.model_json_schema()
def final_score(job,Resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
act as an hr recuiter 
compare the candidate resume with the job description.

Job Description:
{job.model_dump_json(indent = 2)}

Candidate Resume:
{Resume.model_dump_json(indent = 2)}

return json matching this schema
{match_schema}

Give me:
1. Cnadidate name
2. Matching skill
3. missing important skill
4. wheteher experiennce requirement is met
5.overall match percentage from 0 to 100
6. a final verdict or a short summary 
keep the response concise and easy to read."""
    # FIX: the resume was never included in the prompt above


    message = {
        "role" : "user",
        "content" : prompt
    }

    messages = [message]
    response_format = {
        "type" : "json_object"   # FIX: was "json_objct"

    }

    response = client.chat.completions.create(model = model   # FIX: was completion
     ,messages = messages,response_format = response_format)

    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)

def parse_resume(resume_text):
    system_prompt = f"""You are an expert resume parser and HR assistant.

Your task is to extract structured information from the
provided resume text.

Experience may appear under different section headings, such as:
- Experience
- Professional Experience
- Work History
- Employment
- Internships

These sections may all contain relevant experience.

Skills may appear throughout the resume, including the skills
section, work experience, internships, and projects.

Return ONLY valid JSON matching this schema:
{resume_schema}

Important rules:
1. Do not invent information.
2. If a value is not available, return null.
3. If a list has no information, return an empty list.
4. Include internships inside experiences.
5. Extract skills mentioned across the entire resume.
6. Extract each work experience as a separate object.
7. Include company, role, duration, description, and skills used
   whenever this information is available.
8. Extract education, projects, and certifications when present.
9. Use only information supported by the resume.
10. Do not include explanations or Markdown outside the JSON.
"""

    user_prompt = f"""Parse the following resume:

{resume_text}"""

    message_system ={
        "role":"system",
        "content" : system_prompt
    }
    message_user  = {
        "role" : "user",
        "content" : user_prompt

    }
    messages = [message_system, message_user]   # FIX: was messages_system (the job description prompt)
    response_format = {
        "type" : "json_object"
    }
    response = client.chat.completions.create(model=model, messages = messages, response_format =response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume

from pypdf import PdfReader   # FIX: was PdfReadrer
from docx import Document
def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text = text + page_text + "\n"

    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text = text + paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text = text + cell.text + "\n"
    return text


def read_resume_file(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None

import time

resume_folder = Path("resumes")
all_results = []

for file_path in resume_folder.iterdir():
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue

    print("\nProcessing:", file_path.name)

    resume_text = read_resume_file(file_path)
    parsed_resume = parse_resume(resume_text)

    time.sleep(5)

    result = final_score(job, parsed_resume)

    time.sleep(5)

    print("Score:", result.match_score)

    all_results.append({
        "name": parsed_resume.name,
        "score": result.match_score,
        "details": result.details
    })

all_results.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)
print("\n\n========== FINAL CANDIDATE RANKING ==========")

def print_result(rank, candidate):
    """Print one candidate's ranking, score, and evaluation details."""
    print(f"\n{rank}. {candidate['name'] or 'Unknown candidate'}")
    print(f"Score: {candidate['score']}")
    print("Details:", candidate["details"])


for rank, candidate in enumerate(all_results, start=1):
    print_result(rank, candidate)

output_file = Path("resume_results.json")
output_file.write_text(
    json.dumps(all_results, indent=2, ensure_ascii=False),
    encoding="utf-8",
)
print(f"\nResults saved to: {output_file.resolve()}")

# FIX: removed `if __name__ == "__main__": main()` because main() was never defined